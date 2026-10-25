"""Time-range CSV and Excel exports."""

from __future__ import annotations

import asyncio
import csv
import io
from datetime import datetime, timedelta, timezone

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from api.export import RangeError, TelemetryRecorder, iso_ms, parse_range, parse_ts, zone

NOW = datetime(2026, 10, 25, 12, 0, tzinfo=timezone.utc)


def test_parse_range():
    s, e = parse_range(None, None, 15, now=NOW)
    assert (e - s) == timedelta(minutes=15) and e == NOW
    s, e = parse_range("2026-10-25T09:00:00Z", "2026-10-25T16:30:00+05:30", None, now=NOW)
    assert iso_ms(s) == "2026-10-25T09:00:00.000Z" and iso_ms(e) == "2026-10-25T11:00:00.000Z"  # IST → UTC
    assert parse_ts("2026-10-25T09:00:00") == datetime(2026, 10, 25, 9, tzinfo=timezone.utc)  # no zone → UTC
    assert parse_range(None, None, None, now=NOW)[0] == NOW - timedelta(minutes=60)  # default: last hour
    for bad in (("2026-10-25T10:00:00Z", "2026-10-25T09:00:00Z"), ("2026-08-01T00:00:00Z", "2026-10-25T00:00:00Z"),
                ("yesterday", None)):
        with pytest.raises(RangeError):
            parse_range(bad[0], bad[1], None, now=NOW)
    with pytest.raises(RangeError):
        zone("Mars/Olympus")
    assert zone("Asia/Kolkata").key == "Asia/Kolkata"


def test_recorder_samples_every_2_s_and_prunes():
    written, pruned = [], []
    rec = TelemetryRecorder(write=written.extend, prune=pruned.append, retention_h=1)
    snap = {"type": "snapshot", "machine_id": "VMC-01", "ts": "2026-10-25T12:00:00.000Z", "temp_c": 40.0,
            "severity": {"zone": "A"}, "risk": {"rules": "LOW", "ml": {"failure_risk_pct": 3.0}}, "signals": {"tool_wear_min": 120}}
    for _ in range(8):
        rec.observe([snap, {**snap, "machine_id": "LATHE-01"}, {"type": "line"}])
    assert len(written) == 4 and {r["machine_id"] for r in written} == {"VMC-01", "LATHE-01"}
    assert written[0]["zone"] == "A" and written[0]["failure_risk_pct"] == 3.0 and written[0]["tool_wear_min"] == 120
    assert len(pruned) == 1  # at most every 10 min


@pytest.fixture
def client(monkeypatch):
    monkeypatch.setenv("UDYAM_TICKER", "0")
    monkeypatch.setenv("UDYAM_AUTO_DEMO", "0")
    from api.app import app, broadcast_messages, factory_engine
    with TestClient(app) as c:
        t0 = datetime.now(timezone.utc)

        async def go(n):
            for _ in range(n):
                await broadcast_messages(factory_engine.step())
        asyncio.run(go(100))
        yield c, t0


def _rows(text: str) -> list[dict]:
    return list(csv.DictReader(io.StringIO(text)))


def test_csv_sections_honour_the_range(client):
    c, t0 = client
    start = iso_ms(t0 - timedelta(seconds=5))
    end = iso_ms(datetime.now(timezone.utc) + timedelta(seconds=5))
    r = c.get("/api/reports/export.csv", params={"section": "telemetry", "start": start, "end": end})
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/csv")
    assert "attachment" in r.headers["content-disposition"] and "-to-" in r.headers["content-disposition"]
    rows = _rows(r.text)
    assert len(rows) >= 4 * 20 and {"VMC-01", "LATHE-01", "SAW-01", "COMP-01"} <= {x["machine_id"] for x in rows}
    assert all(start <= x["ts"] <= end for x in rows) and rows[0]["zone"]
    # a range in the future has no telemetry
    future = c.get("/api/reports/export.csv", params={"section": "telemetry", "start": "2099-01-01T00:00:00Z",
                                                      "end": "2099-01-01T01:00:00Z"})
    assert future.status_code == 200 and _rows(future.text) == []
    for section in ("events", "alerts", "work_orders", "rl_decisions"):
        assert c.get("/api/reports/export.csv", params={"section": section, "minutes": 5}).status_code == 200
    assert c.get("/api/reports/export.csv", params={"section": "summary"}).status_code == 200  # since reset, unchanged
    bad = c.get("/api/reports/export.csv", params={"section": "telemetry", "start": end, "end": start})
    assert bad.status_code == 422 and bad.json()["detail"]["error"] == "invalid_range"
    assert c.get("/api/reports/export.csv", params={"section": "nope"}).status_code == 422


def test_excel_workbook(client):
    c, t0 = client
    r = c.get("/api/reports/export.xlsx", params={"start": iso_ms(t0 - timedelta(seconds=5)), "minutes": 1,
                                                 "end": iso_ms(datetime.now(timezone.utc) + timedelta(seconds=5))})
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    assert r.headers["content-disposition"].endswith('.xlsx"')
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["About", "Summary", "Telemetry", "Energy & OEE", "Alerts", "Work orders", "Events",
                             "RL decisions", "RL policy evaluation"]
    tel = wb["Telemetry"]
    header = [c.value for c in tel[2]]
    assert header[:3] == ["ts", "ts (Asia/Kolkata)", "machine_id"]
    first = [c.value for c in tel[3]]
    assert isinstance(first[1], datetime)
    assert first[1] == (parse_ts(first[0]) + timedelta(hours=5, minutes=30)).replace(tzinfo=None)  # IST = UTC + 5:30
    about = {row[0].value: row[1].value for row in wb["About"].iter_rows(min_row=2)}
    assert about["rows_telemetry"] == tel.max_row - 2 and about["local_time_zone"] == "Asia/Kolkata"
    assert c.get("/api/reports/export.xlsx", params={"minutes": 5, "tz": "Mars/Olympus"}).status_code == 422
