"""Time-range exports: CSV per section and one Excel workbook (.xlsx) with a sheet per section.

- Range: `start` / `end` (ISO 8601; a time without a zone is read as UTC) or `minutes` back from now.
  Times are stored and compared in UTC; the workbook adds a local-time column (`tz`, default Asia/Kolkata).
- Telemetry comes from the `telemetry` table (one row per machine every 2 s, kept
  UDYAM_TELEMETRY_RETENTION_H hours, default 168), so any range inside the retention can be exported.
- Every sheet says what its numbers are: simulation data unless tagged LIVE, illustrative assumptions
  where they are assumptions.
"""

from __future__ import annotations

import csv
import io
import os
from datetime import datetime, timedelta, timezone
from typing import Any, Iterable
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

MAX_RANGE = timedelta(days=31)
EXCEL_MAX_ROWS = 1_000_000
RECORD_EVERY_TICKS = 4  # 2 s at the 500 ms tick, like the trend charts
RETENTION_H = float(os.environ.get("UDYAM_TELEMETRY_RETENTION_H", "168"))


class RangeError(ValueError):
    """An invalid export range or time zone (HTTP 422)."""


def iso_ms(dt: datetime) -> str:
    """The stored timestamp format: 2026-10-25T10:00:00.000Z."""
    dt = dt.astimezone(timezone.utc)
    return dt.strftime("%Y-%m-%dT%H:%M:%S.") + f"{dt.microsecond // 1000:03d}Z"


def parse_ts(value: str) -> datetime:
    try:
        dt = datetime.fromisoformat(value.strip().replace("Z", "+00:00").replace(" ", "T"))
    except ValueError as err:
        raise RangeError(f"'{value}' is not an ISO 8601 time, e.g. 2026-10-25T10:00:00Z") from err
    return dt.replace(tzinfo=timezone.utc) if dt.tzinfo is None else dt.astimezone(timezone.utc)


def parse_range(start: str | None, end: str | None, minutes: float | None,
                now: datetime | None = None, default_minutes: float = 60.0) -> tuple[datetime, datetime]:
    now = now or datetime.now(timezone.utc)
    end_dt = parse_ts(end) if end else now
    if start:
        start_dt = parse_ts(start)
    else:
        start_dt = end_dt - timedelta(minutes=minutes if minutes is not None else default_minutes)
    if start_dt >= end_dt:
        raise RangeError("start must be before end")
    if end_dt - start_dt > MAX_RANGE:
        raise RangeError("the range can be at most 31 days")
    return start_dt, end_dt


def zone(tz: str) -> ZoneInfo:
    try:
        return ZoneInfo(tz)
    except (ZoneInfoNotFoundError, ValueError) as err:
        raise RangeError(f"unknown time zone '{tz}' (use an IANA name such as Asia/Kolkata or UTC)") from err


def in_range(ts: str | None, start: str, end: str) -> bool:
    return bool(ts) and start <= ts <= end


# ----------------------------------------------------------------------------- telemetry recorder
class TelemetryRecorder:
    """Turns snapshot frames into telemetry rows every 2 s per machine and prunes old rows."""

    def __init__(self, write, prune, retention_h: float = RETENTION_H):
        self.write, self.prune, self.retention_h = write, prune, retention_h
        self._ticks: dict[str, int] = {}
        self._last_prune: datetime | None = None

    def observe(self, msgs: Iterable[dict[str, Any]]) -> None:
        rows = []
        for s in msgs:
            if s.get("type") != "snapshot":
                continue
            mid = s["machine_id"]
            n = self._ticks[mid] = self._ticks.get(mid, 0) + 1
            if n % RECORD_EVERY_TICKS:
                continue
            risk = s.get("risk") or {}
            rows.append({
                "ts": s["ts"], "machine_id": mid, "mode": s.get("mode"), "state": s.get("state"), "rpm": s.get("rpm"),
                "derate_pct": s.get("derate_pct"), "vibration_rms_mm_s": s.get("vibration_rms_mm_s"),
                "zone": (s.get("severity") or {}).get("zone"), "temp_c": s.get("temp_c"), "power_kw": s.get("power_kw"),
                "load_pct": s.get("load_pct"), "energy_kwh": s.get("energy_kwh"),
                "tool_wear_min": (s.get("signals") or {}).get("tool_wear_min"), "risk": risk.get("rules"),
                "failure_risk_pct": (risk.get("ml") or {}).get("failure_risk_pct"),
                "active_scenario": s.get("active_scenario"),
            })
        if rows:
            self.write(rows)
        now = datetime.now(timezone.utc)
        if self._last_prune is None or now - self._last_prune > timedelta(minutes=10):
            self._last_prune = now
            self.prune(iso_ms(now - timedelta(hours=self.retention_h)))


# ----------------------------------------------------------------------------- writers
def csv_text(rows: list[dict[str, Any]], columns: list[str]) -> str:
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=columns, extrasaction="ignore")
    w.writeheader()
    for r in rows:
        w.writerow({k: ("" if r.get(k) is None else r.get(k)) for k in columns})
    return buf.getvalue()


def _excel_value(v: Any) -> Any:
    if isinstance(v, (dict, list, tuple)):
        return str(v)
    return v


def xlsx_bytes(sheets: list[dict[str, Any]], tz: ZoneInfo) -> bytes:
    """sheets: [{name, columns, rows, note?, time_columns?}]. Time columns get a local-time column beside them."""
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook(write_only=True)
    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E79")
    note_font = Font(italic=True, color="555555")
    tz_name = getattr(tz, "key", str(tz))
    for sh in sheets:
        ws = wb.create_sheet(sh["name"][:31])
        time_cols = [c for c in sh.get("time_columns", []) if c in sh["columns"]]
        columns: list[str] = []
        for c in sh["columns"]:
            columns.append(c)
            if c in time_cols:
                columns.append(f"{c} ({tz_name})")
        widths = [max(10, min(48, len(c) + 2)) for c in columns]
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A3" if sh.get("note") else "A2"
        if sh.get("note"):
            cell = WriteOnlyCell(ws, value=sh["note"])
            cell.font = note_font
            ws.append([cell])
        header = []
        for c in columns:
            cell = WriteOnlyCell(ws, value=c)
            cell.font, cell.fill = head_font, head_fill
            cell.alignment = Alignment(vertical="center")
            header.append(cell)
        ws.append(header)
        for r in sh["rows"][:EXCEL_MAX_ROWS]:
            out = []
            for c in sh["columns"]:
                v = r.get(c)
                out.append(_excel_value(v))
                if c in time_cols:
                    local = None
                    if isinstance(v, str) and v:
                        try:
                            local = parse_ts(v).astimezone(tz).replace(tzinfo=None)
                        except RangeError:
                            local = None
                    cell = WriteOnlyCell(ws, value=local)
                    cell.number_format = "yyyy-mm-dd hh:mm:ss"
                    out.append(cell)
            ws.append(out)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
